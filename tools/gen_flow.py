import json, uuid, sys, os
OUT = os.path.join(os.path.dirname(__file__), '..', 'flows')
WFID = '5d0e5f4a-3c2b-4f6e-9a71-0c1d2e3f4a51'
CR_DV = 'wfp_sharedcommondataserviceforapps_a1f3c'
CR_XL = 'wfp_sharedexcelonlinebusiness_b7d2e'
DV_API = '/providers/Microsoft.PowerApps/apis/shared_commondataserviceforapps'
XL_API = '/providers/Microsoft.PowerApps/apis/shared_excelonlinebusiness'

# (logical, excel header, kind, extra)
TABLES = [
 dict(key='areas', plural='wfp_geographicareas', pk='wfp_geographicareaid', xl='tblAreas', sheet='Geographic Area',
      name='Name', fields=[('wfp_code','Code','text',None)]),
 dict(key='units', plural='wfp_units', pk='wfp_unitid', xl='tblUnits', sheet='Unit', name='Name',
      fields=[('wfp_unitcode','Unit code','text',None),('wfp_geographicarea','Geographic area','lookup','areas')]),
 dict(key='positions', plural='wfp_positions', pk='wfp_positionid', xl='tblPositions', sheet='Position', name='Position ID',
      fields=[('wfp_title','Title','text',None),('wfp_unit','Unit','lookup','units'),('wfp_geographicarea','Geographic area','lookup','areas'),
        ('wfp_payplan','Pay plan','choice',['GS','GW','WG']),('wfp_grade','Grade range','text',None),('wfp_topgrade','Full performance grade','int',None),
        ('wfp_employmenttype','Employment type','choice',['Full time','Career seasonal','Seasonal']),
        ('wfp_status','Position status','choice',['Filled','Vacant']),('wfp_incumbent','Incumbent','text',None),
        ('wfp_incumbentemail','Incumbent email','text',None),('wfp_previousincumbent','Previous incumbent','text',None),
        ('wfp_fundingstatus','Funding status','choice',['Approved','Pending','Not approved'])]),
 dict(key='actions', plural='wfp_hiringactions', pk='wfp_hiringactionid', xl='tblHiringActions', sheet='Hiring Action', name='Action ID',
      fields=[('wfp_position','Position','lookup','positions'),('wfp_geographicarea','Geographic area','lookup','areas'),
        ('wfp_stage','Stage','choice',['Draft','PMC review','Cleared to advertise','Announced','Certificate issued','Interviews','Selection approval','Selection approved','Offer']),
        ('wfp_candidate','Selected candidate','text',None),('wfp_proposedgrade','Proposed grade and step','text',None),('wfp_entryonduty','Entry on duty','date',None)]),
 dict(key='steps', plural='wfp_selectionsteps', pk='wfp_selectionstepid', xl='tblApprovalSteps', sheet='Approval Step', name='Step',
      fields=[('wfp_hiringaction','Hiring action','lookup','actions'),('wfp_geographicarea','Geographic area','lookup','areas'),
        ('wfp_sequence','Sequence','int',None),
        ('wfp_approvergroup','Approver group','choice',['Supervisor','Budget','Assistant Chief Operations','Assistant Chief Support Services','Area chief or deputy']),
        ('wfp_decision','Decision','choice',['Waiting','Pending','Approved','Returned']),('wfp_decidedby','Decided by','text',None),
        ('wfp_decidedon','Decided on','date',None),('wfp_comment','Comment','text',None),('wfp_round','Approval round','int',None),
        ('wfp_supervisoremail','Supervisor email','text',None),('wfp_positiontext','Position text','text',None)]),
 dict(key='details', plural='wfp_details', pk='wfp_detailid', xl='tblDetails', sheet='Detail', name='Detail ID',
      fields=[('wfp_homeposition','Home position','lookup','positions'),('wfp_detailedto','Detailed to','lookup','positions'),
        ('wfp_geographicarea','Geographic area','lookup','areas'),('wfp_enteredby','Entered by','text',None)]),
 dict(key='roles', plural='wfp_roleassignments', pk='wfp_roleassignmentid', xl='tblRoles', sheet='Role Assignment', name='Person',
      fields=[('wfp_useremail','User email','text',None),
        ('wfp_role','Role','choice',['Budget','Area fire chief','Deputy chief','Assistant Chief Operations','Assistant Chief Support Services','PMC member','Data administrator']),
        ('wfp_status','Assignment status','choice',['Active','Ended']),('wfp_geographicarea','Geographic area','lookup','areas'),
        ('wfp_allareas','All areas','bool',None),('wfp_startdate','Start date','date',None)]),
]
T = {t['key']: t for t in TABLES}
LOOKUP_SRC = {}  # key -> action name of "list after" for that table

def dv(op, params, runafter=None, extra=None):
    a = {"type":"OpenApiConnection","inputs":{"host":{"connectionName":"shared_commondataserviceforapps","operationId":op,"apiId":DV_API},"parameters":params,"authentication":"@parameters('$authentication')"},"runAfter":runafter or {}}
    if extra: a.update(extra)
    return a

def val_expr(item, f):
    lg, hdr, kind, extra = f
    cell = "%s?['%s']" % (item, hdr)
    if kind == 'text': return "@%s" % cell
    if kind == 'int': return "@if(empty(string(%s)), null, int(%s))" % (cell, cell)
    if kind == 'date': return "@if(empty(string(%s)), null, formatDateTime(%s, 'yyyy-MM-dd'))" % (cell, cell)
    if kind == 'bool':
        s = "toLower(string(%s))" % cell
        return "@or(equals(%s,'true'), equals(%s,'yes'), equals(%s,'1'))" % (s, s, s)
    if kind == 'choice':
        e = "null"
        for i, lab in reversed(list(enumerate(extra))):
            e = "if(equals(%s,'%s'), %d, %s)" % (cell, lab.replace("'", "''"), 100000000 + i, e)
        return "@" + e

def build():
    actions = {}
    prev = None
    def chain(name, act):
        nonlocal prev
        act['runAfter'] = {prev: ['Succeeded']} if prev else {}
        actions[name] = act
        prev = name
    for t in TABLES:
        k = t['key']
        chain('Read_' + k, {"type":"OpenApiConnection","inputs":{"host":{"connectionName":"shared_excelonlinebusiness","operationId":"GetItems","apiId":XL_API},
            "parameters":{"source":"me","drive":"me","file":"@triggerBody()?['text']","table":t['xl'],"dateTimeFormat":"ISO 8601"},"authentication":"@parameters('$authentication')"}})
        chain('Existing_' + k, dv('ListRecords', {"entityName":t['plural'],"$select":"%s,wfp_name" % t['pk'],"$top":5000}))
        fe = 'Each_' + k
        body = {}
        # lookup filters
        inner = {}
        lastinner = None
        def ichain(n, a):
            nonlocal lastinner
            a['runAfter'] = {lastinner: ['Succeeded']} if lastinner else {}
            inner[n] = a; lastinner = n
        ichain('Self_' + k, {"type":"Query","inputs":{"from":"@body('Existing_%s')?['value']" % k,"where":"@equals(item()?['wfp_name'], items('%s')?['%s'])" % (fe, t['name'])}})
        item = "items('%s')" % fe
        fields = {}
        fields['item/wfp_name'] = "@%s?['%s']" % (item, t['name'])
        for f in t['fields']:
            lg, hdr, kind, extra = f
            if kind == 'lookup':
                fn = 'Look_%s_%s' % (k, lg)
                ichain(fn, {"type":"Query","inputs":{"from":"@body('After_%s')?['value']" % extra,"where":"@equals(item()?['wfp_name'], %s?['%s'])" % (item, hdr)}})
                ptable = T[extra]
                fields['item/%s@odata.bind' % lg] = "@if(empty(body('%s')), null, concat('/%s(', first(body('%s'))?['%s'], ')'))" % (fn, ptable['plural'], fn, ptable['pk'])
            else:
                fields['item/' + lg] = val_expr(item, f)
        upd = dict(fields); upd['entityName'] = t['plural']; upd['recordId'] = "@first(body('Self_%s'))?['%s']" % (k, t['pk'])
        cre = dict(fields); cre['entityName'] = t['plural']
        ichain('Upsert_' + k, {"type":"If","expression":{"equals":["@length(body('Self_%s'))" % k, 0]},
            "actions":{"Create_"+k: dv('CreateRecord', cre)},
            "else":{"actions":{"Update_"+k: dv('UpdateRecord', upd)}}})
        chain(fe, {"type":"Foreach","foreach":"@body('Read_%s')?['value']" % k,"actions":inner,"runtimeConfiguration":{"concurrency":{"repetitions":1}}})
        chain('After_' + k, dv('ListRecords', {"entityName":t['plural'],"$select":"%s,wfp_name" % t['pk'],"$top":5000}))
        if k == 'positions':
            # second pass: supervisors
            inner2 = {}; li = None
            def i2(n, a):
                nonlocal li
                a['runAfter'] = {li: ['Succeeded']} if li else {}
                inner2[n] = a; li = n
            it = "items('Each_supervisors')"
            i2('SelfRow', {"type":"Query","inputs":{"from":"@body('After_positions')?['value']","where":"@equals(item()?['wfp_name'], %s?['Position ID'])" % it}})
            i2('SupRow', {"type":"Query","inputs":{"from":"@body('After_positions')?['value']","where":"@equals(item()?['wfp_name'], %s?['Supervisor position'])" % it}})
            i2('SetSup', {"type":"If","expression":{"and":[{"greater":["@length(body('SelfRow'))",0]},{"greater":["@length(body('SupRow'))",0]}]},
                "actions":{"UpdateSup": dv('UpdateRecord', {"entityName":"wfp_positions","recordId":"@first(body('SelfRow'))?['wfp_positionid']",
                    "item/wfp_supervisor@odata.bind":"@concat('/wfp_positions(', first(body('SupRow'))?['wfp_positionid'], ')')"})}})
            chain('Each_supervisors', {"type":"Foreach","foreach":"@body('Read_positions')?['value']","actions":inner2,"runtimeConfiguration":{"concurrency":{"repetitions":1}}})
    return actions

definition = {
 "$schema":"https://schema.management.azure.com/providers/Microsoft.Logic/schemas/2016-06-01/workflowdefinition.json#",
 "contentVersion":"1.0.0.0",
 "parameters":{"$connections":{"defaultValue":{},"type":"Object"},"$authentication":{"defaultValue":{},"type":"SecureObject"}},
 "triggers":{"manual":{"metadata":{"operationMetadataId":str(uuid.uuid4())},"type":"Request","kind":"Button",
   "inputs":{"schema":{"type":"object","properties":{"text":{"title":"Excel file path in OneDrive","type":"string","x-ms-dynamically-added":True,
     "description":"For example /WorkforcePlanner_Import.xlsx","x-ms-content-hint":"TEXT"}},"required":["text"]}}}},
 "actions": build()}
flow = {"properties":{"connectionReferences":{
   "shared_commondataserviceforapps":{"runtimeSource":"embedded","connection":{"connectionReferenceLogicalName":CR_DV},"api":{"name":"shared_commondataserviceforapps"}},
   "shared_excelonlinebusiness":{"runtimeSource":"embedded","connection":{"connectionReferenceLogicalName":CR_XL},"api":{"name":"shared_excelonlinebusiness"}}},
  "definition":definition,"templateName":None},"schemaVersion":"1.0.0.0"}
json.dump(flow, open(os.path.join(OUT,'WorkforcePlannerImport-%s.json' % WFID.upper()),'w'), indent=2)

# --- sample workbook with Excel tables
import openpyxl
from openpyxl.worksheet.table import Table, TableStyleInfo
seed = json.load(open(os.path.join(os.path.dirname(__file__),'..','seed_data.json')))
wb = openpyxl.Workbook(); wb.remove(wb.active)
sup = {a:b for a,b in seed['_supervisors']}
def nm(v): return v[1] if isinstance(v, list) else v
for t in TABLES:
    key = {'areas':'wfp_geographicarea','units':'wfp_unit','positions':'wfp_position','actions':'wfp_hiringaction','steps':'wfp_selectionstep','details':'wfp_detail','roles':'wfp_roleassignment'}[t['key']]
    ws = wb.create_sheet(t['sheet'])
    hdr = [t['name']] + [f[1] for f in t['fields']] + (['Supervisor position'] if t['key']=='positions' else [])
    ws.append(hdr)
    for r in seed[key]:
        row = [r['wfp_name']]
        for lg, h, kind, extra in t['fields']:
            v = r.get(lg)
            if h == 'Position text' : v = r.get('wfp_positiontext')
            if h == 'Selected candidate': v = r.get('wfp_candidate')
            if kind == 'lookup': v = nm(v) if v else ''
            row.append('' if v is None else v)
        if t['key']=='positions': row.append(sup.get(r['wfp_name'],''))
        ws.append(row)
    ref = "A1:%s%d" % (openpyxl.utils.get_column_letter(len(hdr)), max(2, len(seed[key])+1))
    tab = Table(displayName=t['xl'], ref=ref); tab.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(tab)
    for i in range(len(hdr)): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width = 22
wb.save(os.path.join(OUT,'WorkforcePlanner_Import_Sample.xlsx'))
print('ok', len(json.dumps(flow)))
